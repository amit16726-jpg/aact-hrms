"""Makes real report files through the app (fake Supabase as in screens.py) and checks them.

Usage: python3 e2e/report_files.py <dist> <out-dir>
Downloads Excel and PDF of P01 (payroll register) and W02 (employee list, 120 rows) in English and Arabic and the
PowerPoint board pack, then opens each file: workbook sheets, header, totals and number formats; PDF page count and size;
slides and charts. Prints FILES OK or the problems.
"""
import json, sys, threading, time, http.server, functools, pathlib, zipfile
from playwright.sync_api import sync_playwright

DIST, OUT = sys.argv[1], pathlib.Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)
src = pathlib.Path(__file__).with_name('screens.py').read_text(encoding='utf-8')
# reuse the fixtures and the server from screens.py without running its screenshot loop
head = src.split('problems = []')[0].replace("srv = http.server.ThreadingHTTPServer(('127.0.0.1', 4310)", "srv = http.server.ThreadingHTTPServer(('127.0.0.1', 4311)")
sys.argv = [sys.argv[0], DIST, str(OUT / '_unused')]
g = {'__file__': str(pathlib.Path(__file__).with_name('screens.py').resolve())}
exec(compile(head, 'screens.py', 'exec'), g)
RPC, ME, SESSION = g['RPC'], g['ME'], g['SESSION']

problems = []
files = {}
with sync_playwright() as p:
    br = p.chromium.launch()
    for lang in ('en', 'ar'):
        ctx = br.new_context(viewport={'width': 1440, 'height': 900}, accept_downloads=True, service_workers='block')
        ctx.add_init_script(f"localStorage.setItem('sb-mockref-auth-token', {json.dumps(SESSION)}); localStorage.setItem('aact-hrms-lang', '{lang}');")
        def handle(route):
            fn = route.request.url.split('/rpc/')[-1].split('?')[0]
            body = RPC.get(fn)
            if fn == 'hr_run_report':
                body = RPC['hr_run_report'].get((route.request.post_data_json or {}).get('p_code', 'P01'))
            if fn == 'hr_me':
                body = {**ME, 'language': lang}
            route.fulfill(status=200, content_type='application/json', body=json.dumps(body if body is not None else {}))
        ctx.route('**/rest/v1/rpc/**', handle)
        ctx.route('**/auth/v1/**', lambda r: r.fulfill(status=200, content_type='application/json', body='{}'))
        page = ctx.new_page()
        page.on('pageerror', lambda e: problems.append(f'pageerror {e}'))
        for code in ('P01', 'W02'):
            page.goto(f'http://127.0.0.1:4311/reports?code={code}&run=1')
            page.wait_for_load_state('networkidle')
            page.wait_for_selector('.rp-table, .empty', timeout=15000)
            for kind, label in (('xlsx', 'Excel'), ('pdf', 'PDF')):
                with page.expect_download(timeout=120000) as d:
                    page.get_by_role('button', name=label, exact=True).click()
                path = OUT / f'{code}-{lang}.{kind}'
                d.value.save_as(path)
                files[f'{code}-{lang}.{kind}'] = path
        page.goto('http://127.0.0.1:4311/reports?tab=board')
        page.wait_for_load_state('networkidle')
        with page.expect_download(timeout=120000) as d:
            page.locator('button:has(svg.lucide-presentation)').last.click()
        path = OUT / f'board-{lang}.pptx'
        d.value.save_as(path)
        files[f'board-{lang}.pptx'] = path
        ctx.close()
    br.close()

# ----- check the files
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation
for name, path in files.items():
    size = path.stat().st_size
    if size < 1000: problems.append(f'{name}: only {size} bytes'); continue
    if name.endswith('.xlsx'):
        wb = load_workbook(path)
        ws = wb.active
        hdr = [c.value for c in ws[6]]
        if not hdr or hdr[0] is None: problems.append(f'{name}: no header row')
        data_rows = ws.max_row - 7
        exp = 120 if name.startswith('W02') else len(RPC['hr_run_report']['P01']['rows'])
        if name.startswith('P01'):
            last = [c.value for c in ws[ws.max_row - 1]]
            if not any(isinstance(v, (int, float)) for v in last): problems.append(f'{name}: no totals row')
            if ws.cell(row=7, column=7).number_format != '#,##0.000': problems.append(f'{name}: money format missing ({ws.cell(row=7, column=7).number_format})')
        if ws.auto_filter.ref is None: problems.append(f'{name}: no filter')
        if (ws.sheet_view.rightToLeft or False) != name.endswith('-ar.xlsx'): problems.append(f'{name}: wrong sheet direction')
        print(f'{name}: {size} bytes, {ws.max_row} rows, header {hdr[:4]}')
    elif name.endswith('.pdf'):
        r = PdfReader(str(path))
        n = len(r.pages)
        if name.startswith('W02') and n < 3: problems.append(f'{name}: expected several pages for 120 rows, got {n}')
        print(f'{name}: {size} bytes, {n} pages, {float(r.pages[0].mediabox.width):.0f}x{float(r.pages[0].mediabox.height):.0f} pt')
    elif name.endswith('.pptx'):
        prs = Presentation(str(path))
        charts = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
        if len(prs.slides) < 8 or charts < 6: problems.append(f'{name}: {len(prs.slides)} slides, {charts} charts')
        print(f'{name}: {size} bytes, {len(prs.slides)} slides, {charts} charts')
print('\n'.join(problems) if problems else 'FILES OK')
sys.exit(1 if problems else 0)
